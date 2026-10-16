import hashlib
import io
import re
import zipfile
from pathlib import Path

import openpyxl
import pytest

import categories as C
import reports

ROWS = [
    {"id": 1, "item": "10 ta paypoq oldim", "amount": 13000, "cat": "daily.clothing", "ts": "2026-10-12T18:10:00"},
    {"id": 2, "item": "taksi", "amount": 25000, "cat": "transport.fare", "ts": "2026-10-13T09:00:00"},
    {"id": 3, "item": "taksi", "amount": 15000, "cat": "transport.fare", "ts": "2026-10-14T09:00:00"},
    {"id": 4, "item": "???", "amount": 7000, "cat": None, "ts": "2026-10-15T12:00:00"},
    {"id": 5, "item": "ijara", "amount": 2_500_000, "cat": "home.rent", "ts": "2026-09-01T12:00:00"},
    {"id": 6, "item": "old year", "amount": 999, "cat": "home.rent", "ts": "2025-12-31T12:00:00"},
]
TEMPLATE = reports.TEMPLATE


def _load(lang, **kw):
    return openpyxl.load_workbook(reports.build_excel(ROWS, "T", lang, "p", **kw), data_only=True)


def test_template_file_is_never_modified():
    before = hashlib.md5(TEMPLATE.read_bytes()).hexdigest()
    reports.build_excel(ROWS, "T", "uz")
    assert hashlib.md5(TEMPLATE.read_bytes()).hexdigest() == before


@pytest.mark.parametrize("lang", ["uz", "ru", "en"])
def test_excel_filled_and_localized(lang):
    wb = _load(lang, year=2026)
    ws = wb["Budget"]
    rows_by_name = {ws.cell(r, 1).value: r for r in range(1, 105)}
    # Oct = column K (11th), Sep = J
    clothing = C.sub_name(lang, "daily.clothing"); fare = C.sub_name(lang, "transport.fare")
    assert ws[f"K{rows_by_name[clothing]}"].value == 13000
    assert ws[f"K{rows_by_name[fare]}"].value == 40000                      # 25k + 15k summed in the month
    assert ws[f"J{rows_by_name[C.sub_name(lang, 'home.rent')]}"].value == 2_500_000
    assert ws["K101"].value == 7000                                         # unknown category -> misc "Other"
    # the 2025 expense must not leak into the 2026 sheet
    assert ws[f"N{rows_by_name[C.sub_name(lang, 'home.rent')]}"].value == 2_500_000
    assert ws["A1"].value == reports.t(lang, "budget_title")
    assert ws["B3"].value == reports.t(lang, "months")[0] and ws["M3"].value == reports.t(lang, "months")[11]
    total_word = reports.t(lang, "tbl_total")
    assert ws["A17"].value == f"{total_word} {C.GROUPS[0][C.LANG_IDX[lang]]}"
    assert ws["N103"].value == 7000 and ws["N27"].value == 40000


@pytest.mark.parametrize("lang", ["uz", "ru"])
def test_tables_consistent_for_excel(lang):
    """Excel 'repairs' a file if table column names differ from header cells or formulas point nowhere."""
    wb = openpyxl.load_workbook(reports.build_excel(ROWS, "T", lang))
    ws = wb["Budget"]
    for tb in ws.tables.values():
        c0, r0, c1, r1 = openpyxl.utils.cell.range_boundaries(tb.ref)
        names = [c.name for c in tb.tableColumns]
        assert names == [ws.cell(r0, c).value for c in range(c0, c1 + 1)]
        assert len(set(names)) == len(names)
        formulas = [ws.cell(r1, c).value for c in range(c0, c1 + 1)] + \
                   [c.totalsRowFormula.attr_text for c in tb.tableColumns if c.totalsRowFormula is not None]
        for f in formulas:
            for tok in re.findall(r"\[([^\[\]#]+)\]", str(f)):
                assert tok in names, (tb.name, tok)
            assert "Total " not in str(f) or lang == "en"
    assert all(ws.cell(r, 1).value != "Other" for r in range(1, 105))


def test_cached_values_match_libreoffice_recalc(tmp_path):
    import shutil, subprocess
    if not shutil.which("soffice"):
        pytest.skip("LibreOffice not installed")
    f = tmp_path / "r.xlsx"
    f.write_bytes(reports.build_excel(ROWS, "T", "uz", year=2026).getvalue())
    out = tmp_path / "lo"
    subprocess.run(["soffice", "--headless", "--convert-to", "xlsx", "--outdir", str(out), str(f)],
                   check=True, capture_output=True, timeout=120)
    mine = openpyxl.load_workbook(f, data_only=True)["Budget"]
    lo = openpyxl.load_workbook(out / "r.xlsx", data_only=True)["Budget"]
    checked = 0
    for r in range(3, 104):
        for c in range(1, 16):
            a, b = mine.cell(r, c).value, lo.cell(r, c).value
            if isinstance(a, (int, float)) or isinstance(b, (int, float)):
                assert abs((a or 0) - (b or 0)) < 1e-6, (r, c, a, b)
                checked += 1
            elif a is not None and b is not None:
                assert a == b, (r, c, a, b)
    assert checked > 250


def test_template_formatting_preserved():
    orig = openpyxl.load_workbook(TEMPLATE)["Budget"]
    new = openpyxl.load_workbook(reports.build_excel(ROWS, "T", "uz"))["Budget"]
    for coord in ("A3", "B3", "B4", "N4", "A17", "B17", "A19"):
        a, b = orig[coord], new[coord]
        assert (a.font.name, a.font.sz, a.font.b, a.number_format) == (b.font.name, b.font.sz, b.font.b, b.number_format)
    assert sorted(t.ref for t in orig.tables.values()) == sorted(t.ref for t in new.tables.values())
    assert new.print_area == orig.print_area
    z = zipfile.ZipFile(reports.build_excel(ROWS, "T", "uz"))
    assert sum(n.startswith("xl/media/") for n in z.namelist()) >= 3     # template logos kept


def test_other_formats_and_empty_rows():
    for lang in ("uz", "ru", "en"):
        assert "%" in reports.build_text(ROWS, "T", lang)
        assert reports.build_pdf(ROWS, "T", lang).read(4) == b"%PDF"
        assert reports.build_image(ROWS, "T", lang).read(4) == b"\x89PNG"
        assert openpyxl.load_workbook(reports.build_excel([], "T", lang)).sheetnames[0] == "Budget"
        assert reports.build_pdf([], "T", lang).read(4) == b"%PDF"


def test_long_period_aggregates_by_month():
    rows = [{"id": i, "item": "x", "amount": 10, "cat": None, "ts": f"2026-{1 + i // 30:02d}-{1 + i % 28:02d}T10:00:00"}
            for i in range(300)]
    s = reports.summarize(rows, "uz")
    assert s["series_label"] == "by_month" and len(s["series"]) <= 12
