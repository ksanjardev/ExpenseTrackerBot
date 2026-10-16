"""Builds text / image / PDF / Excel reports from a list of expense dicts, in uz / ru / en.

The Excel report is the ORIGINAL Vertex42 template (assets/personal-budget-template.xlsx) opened with
openpyxl, localized in place and filled with the user's monthly sums - formatting, tables, formulas and
print setup are the template's own. A Dashboard and a Details sheet are added after it.
"""
import io
import os
import re
import threading
import zipfile
from collections import OrderedDict
from pathlib import Path
from xml.sax.saxutils import escape

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import range_boundaries
from openpyxl.workbook.properties import CalcProperties
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

import categories as C
from expense_parser import fmt_money
from i18n import HELP, t

TEMPLATE = Path(__file__).parent / "assets" / "personal-budget-template.xlsx"

# DejaVu ships with matplotlib -> same file on every OS, covers Latin (incl. Uzbek ʻ ʼ) and Cyrillic
_FONT_DIR = os.path.join(matplotlib.get_data_path(), "fonts", "ttf")
pdfmetrics.registerFont(TTFont("DejaVu", os.path.join(_FONT_DIR, "DejaVuSans.ttf")))
pdfmetrics.registerFont(TTFont("DejaVu-Bold", os.path.join(_FONT_DIR, "DejaVuSans-Bold.ttf")))
pdfmetrics.registerFontFamily("DejaVu", normal="DejaVu", bold="DejaVu-Bold", italic="DejaVu", boldItalic="DejaVu-Bold")

BLUE, ORANGE = "#1F4E78", "#E8833A"
_PLT_LOCK = threading.Lock()          # pyplot keeps global state; reports run in worker threads


def _day(r): return r["ts"][:10]
def _time(r): return r["ts"][11:16]


# ----------------------------------------------------------------------------- summary
def summarize(rows, lang):
    total = sum(r["amount"] for r in rows)
    groups = OrderedDict((g[0], 0) for g in C.GROUPS)
    daily, items = OrderedDict(), {}
    rows = sorted(rows, key=lambda r: r['ts'])
    for r in rows:
        g, _ = C.split_cat(r.get("cat"))
        groups[g] += r["amount"]
        daily[_day(r)] = daily.get(_day(r), 0) + r["amount"]
        key = r["item"].strip().lower()
        items[key] = items.get(key, 0) + r["amount"]
    by_group = sorted(((C.group_name(lang, g).capitalize(), v) for g, v in groups.items() if v),
                      key=lambda kv: kv[1], reverse=True)
    top = sorted(items.items(), key=lambda kv: kv[1], reverse=True)
    # long periods -> aggregate by month so charts stay readable
    if len(daily) > 62:
        series, label = OrderedDict(), "by_month"
        for d, v in daily.items():
            series[d[:7]] = series.get(d[:7], 0) + v
    else:
        series, label = daily, "by_day"
    return dict(total=total, count=len(rows), by_group=by_group, daily=daily, top=top, series=series,
                series_label=label, days=max(len(daily), 1), biggest=max(rows, key=lambda r: r["amount"]) if rows else None)


def pct(v, total):
    if not total:
        return "0%"
    p = round(v * 100 / total)
    return "<1%" if p == 0 and v > 0 else f"{p}%"


def _money(lang, n):
    return f"{fmt_money(n)} {t(lang, 'currency')}"


# ----------------------------------------------------------------------------- text
def build_text(rows, title, lang):
    if not rows:
        return f"{title}\n\n{t(lang, 'no_data_short')}"
    s = summarize(rows, lang)
    L = [f"📊 {title}", f"💰 {t(lang, 'total')}: {_money(lang, s['total'])} ({s['count']})", "",
         f"📂 {t(lang, 'by_cat')}:"]
    L += [f"• {k} — {fmt_money(v)} ({pct(v, s['total'])})" for k, v in s["by_group"]]
    L += ["", f"🏆 {t(lang, 'top_items')}:"]
    L += [f"• {k} — {fmt_money(v)}" for k, v in s["top"][:6]]
    L += ["", f"📅 {t(lang, s['series_label'])}:"]
    L += [f"• {d} — {fmt_money(v)}" for d, v in list(s["series"].items())[-31:]]
    return "\n".join(L)[:4000]


# ----------------------------------------------------------------------------- image
def build_image(rows, title, lang):
    with _PLT_LOCK:
        return _build_image(rows, title, lang)


def _build_image(rows, title, lang):
    s = summarize(rows, lang)
    fig, (a1, a2) = plt.subplots(2, 1, figsize=(9, 9.5), gridspec_kw={"height_ratios": [1, 1.1]})
    fig.suptitle(f"{title}\n{t(lang, 'total')}: {_money(lang, s['total'])}", fontsize=14, fontweight="bold")
    keys = list(s["series"].keys())
    vals = list(s["series"].values())
    a1.bar([k[5:] if len(k) == 10 else k for k in keys], vals, color=BLUE)
    a1.set_title(t(lang, s["series_label"]))
    a1.tick_params(axis="x", rotation=45)
    a1.yaxis.set_major_formatter(lambda x, _: fmt_money(x))
    if len(vals) <= 16:
        for i, v in enumerate(vals):
            a1.text(i, v, fmt_money(v), ha="center", va="bottom", fontsize=8)
    g = s["by_group"][:10][::-1]
    a2.barh([k[:28] for k, _ in g], [v for _, v in g], color=ORANGE)
    a2.set_title(t(lang, "by_cat"))
    for i, (_, v) in enumerate(g):
        a2.text(v, i, f" {fmt_money(v)} ({pct(v, s['total'])})", va="center", fontsize=8)
    a2.margins(x=0.28)
    a2.xaxis.set_major_formatter(lambda x, _: fmt_money(x))
    for a in (a1, a2):
        a.spines[["top", "right"]].set_visible(False)
    fig.tight_layout(rect=(0, 0, 1, 0.94))
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=130)
    plt.close(fig)
    buf.seek(0)
    return buf


# ----------------------------------------------------------------------------- pdf
def build_pdf(rows, title, lang):
    s = summarize(rows, lang)
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, title=title, leftMargin=36, rightMargin=36, topMargin=36,
                            bottomMargin=36)
    ss = getSampleStyleSheet()
    h1 = ParagraphStyle("h1", parent=ss["Title"], fontName="DejaVu-Bold", fontSize=17)
    h2 = ParagraphStyle("h2", parent=ss["Heading2"], fontName="DejaVu-Bold", fontSize=12)
    body = ParagraphStyle("b", parent=ss["Normal"], fontName="DejaVu", fontSize=10.5)

    def table(data, widths):
        tb = Table(data, colWidths=widths, repeatRows=1)
        tb.setStyle(TableStyle([
            ("FONTNAME", (0, 0), (-1, -1), "DejaVu"), ("FONTSIZE", (0, 0), (-1, -1), 8.5),
            ("FONTNAME", (0, 0), (-1, 0), "DejaVu-Bold"),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(BLUE)), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F2F2F2")]),
            ("ALIGN", (-1, 0), (-1, -1), "RIGHT"), ("GRID", (0, 0), (-1, -1), 0.25, colors.lightgrey),
        ]))
        return tb

    story = [Paragraph(escape(title), h1),
             Paragraph(f"{t(lang, 'total')}: <b>{_money(lang, s['total'])}</b> &nbsp;|&nbsp; "
                       f"{t(lang, 'count')}: <b>{s['count']}</b>", body), Spacer(1, 10)]
    if rows:
        story += [Paragraph(t(lang, "by_cat"), h2),
                  table([[t(lang, "category"), t(lang, "amount"), "%"]] +
                        [[k, fmt_money(v), pct(v, s["total"])] for k, v in s["by_group"]],
                        [300, 150, 73]), Spacer(1, 10)]
        story += [Paragraph(t(lang, "top_items"), h2),
                  table([[t(lang, "item"), t(lang, "amount")]] + [[k[:60], fmt_money(v)] for k, v in s["top"][:10]],
                        [370, 153]), Spacer(1, 10)]
        story += [Paragraph(t(lang, s["series_label"]), h2),
                  table([[t(lang, "date"), t(lang, "amount")]] + [[d, fmt_money(v)] for d, v in s["series"].items()],
                        [370, 153]), Spacer(1, 10)]
        story += [Paragraph(t(lang, "all_items"), h2),
                  table([[t(lang, "date"), t(lang, "time"), t(lang, "item"), t(lang, "category"), t(lang, "amount")]] +
                        [[_day(r), _time(r), r["item"][:34], C.group_name(lang, C.split_cat(r.get("cat"))[0]).capitalize()[:22],
                          fmt_money(r["amount"])] for r in rows],
                        [58, 40, 184, 128, 113])]
    else:
        story.append(Paragraph(t(lang, "no_data_short"), body))
    doc.build(story)
    buf.seek(0)
    return buf


# ----------------------------------------------------------------------------- excel (template based)
_TOKEN = re.compile(r"\[([^\[\]]+)\]")


def _loc_formula(f: str, tokmap: dict, total_word: str) -> str:
    """Rewrite a structured-reference formula after the table's column names were localized."""
    f = f.replace('"Total "', f'"{total_word} "')
    return _TOKEN.sub(lambda m: "[" + tokmap.get(m.group(1), m.group(1)) + "]", f)


def _localize_and_fill(wb, rows, lang, year):
    """Edits the Budget sheet in place. Returns {coord: cached_value} for every formula cell."""
    ws = wb["Budget"]
    tables = sorted(ws.tables.values(), key=lambda tb: range_boundaries(tb.ref)[1])
    if len(tables) != len(C.GROUPS):
        raise RuntimeError(f"template has {len(tables)} tables, expected {len(C.GROUPS)}")

    # month sums per (group, sub)
    sums = {}
    for r in rows:
        if r["ts"][:4] != str(year):
            continue
        g, s = C.split_cat(r.get("cat"))
        if not C.valid_cat(f"{g}.{s}"):
            g, s = C.split_cat(None)
        sums[(g, s, int(r["ts"][5:7]) - 1)] = sums.get((g, s, int(r["ts"][5:7]) - 1), 0) + r["amount"]

    months = t(lang, "months")
    total_w, avg_w = t(lang, "tbl_total"), t(lang, "tbl_avg")
    cache = {}
    ws["A1"] = t(lang, "budget_title")

    for tb, grp in zip(tables, C.GROUPS):
        gkey, subs = grp[0], grp[4]
        min_col, hdr, max_col, last = range_boundaries(tb.ref)
        data_rows = list(range(hdr + 1, last))
        if len(data_rows) != len(subs):
            raise RuntimeError(f"{tb.name}: {len(data_rows)} rows in template vs {len(subs)} categories")

        old = [c.name for c in tb.tableColumns]
        new = [grp[C.LANG_IDX[lang]]] + list(months) + [total_w, avg_w]
        tokmap = dict(zip(old, new))
        for col, nm in zip(tb.tableColumns, new):               # keep table metadata == header cells
            col.name = nm
        for i, nm in enumerate(new):
            ws.cell(hdr, 1 + i).value = nm

        col_tot = [0] * 12
        for sub, r in zip(subs, data_rows):
            ws.cell(r, 1).value = sub[C.LANG_IDX[lang]]
            row_total = 0
            for m in range(12):
                v = sums.get((gkey, sub[0], m), 0)
                if v:
                    ws.cell(r, 2 + m).value = v
                row_total += v
                col_tot[m] += v
            cache[f"N{r}"] = row_total
            cache[f"O{r}"] = row_total / 12

        # totals row: template formulas, re-pointed at the renamed columns
        for c in range(1, max_col + 1):
            cell = ws.cell(last, c)
            if isinstance(cell.value, str) and cell.value.startswith("="):
                cell.value = "=" + _loc_formula(cell.value[1:], tokmap, total_w)
        for col in tb.tableColumns:
            if col.totalsRowFormula is not None and col.totalsRowFormula.attr_text:
                col.totalsRowFormula.attr_text = _loc_formula(col.totalsRowFormula.attr_text, tokmap, total_w)
        cache[f"A{last}"] = f"{total_w} {new[0]}"
        for m in range(12):
            cache[f"{get_column_letter(2 + m)}{last}"] = col_tot[m]
        cache[f"N{last}"] = sum(col_tot)
        cache[f"O{last}"] = sum(col_tot) / 12

    # amounts are in so'm: the template's narrow columns would show ####
    for d in ws.column_dimensions.values():
        if d.min == 1:
            d.width = 34
        elif d.min is not None and 2 <= d.min <= 13:
            d.width = 11.5
        elif d.min in (14, 15):
            d.width = 13

    # Help sheet text
    if lang in HELP:
        for coord, text in HELP[lang].items():
            wb["Help"][coord].value = text
    return cache


def _style_header(cells):
    for c in cells:
        c.font = Font(name="Arial", bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", start_color=BLUE.lstrip("#"))
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _add_dashboard_and_details(wb, rows, lang, title, period_text):
    s = summarize(rows, lang)
    base, bold = Font(name="Arial", size=10), Font(name="Arial", size=10, bold=True)

    # ---- details sheet (raw data, handy for filters)
    det = wb.create_sheet(t(lang, "details"), 2)   # Budget, Dashboard, Details, Help, ©
    det.append([t(lang, k) for k in ("date", "time", "item", "category", "subcategory", "amount")])
    _style_header(det[1])
    for r in rows:
        g, _ = C.split_cat(r.get("cat"))
        det.append([_day(r), _time(r), r["item"], C.group_name(lang, g).capitalize(),
                    C.sub_name(lang, r.get("cat") or C.FALLBACK_CAT), r["amount"]])
    for row in det.iter_rows(min_row=2):
        for c in row:
            c.font = base
        row[5].number_format = "#,##0"
    for col, w in zip("ABCDEF", (12, 8, 38, 28, 28, 16)):
        det.column_dimensions[col].width = w
    det.freeze_panes = "A2"
    det.auto_filter.ref = f"A1:F{max(len(rows) + 1, 2)}"

    # ---- dashboard (static values: opens correctly even in viewers that don't recalculate)
    d = wb.create_sheet(t(lang, "dashboard"), 1)
    d.sheet_view.showGridLines = False
    d["A1"], d["A1"].font = title, Font(name="Arial", bold=True, size=14, color=BLUE.lstrip("#"))
    d["A2"], d["A2"].font = f"{t(lang, 'period')}: {period_text}", Font(name="Arial", size=9, color="7F7F7F")
    kpis = [(t(lang, "total"), s["total"]), (t(lang, "count"), s["count"]),
            (t(lang, "avg_day"), round(s["total"] / s["days"])),
            (t(lang, "biggest"), s["biggest"]["amount"] if s["biggest"] else 0)]
    for i, (k, v) in enumerate(kpis):
        d.cell(4 + i, 1, k).font = bold
        c = d.cell(4 + i, 2, v)
        c.font, c.number_format = base, "#,##0"

    r0 = 10
    d.cell(r0, 1, t(lang, "category")), d.cell(r0, 2, t(lang, "amount"))
    _style_header([d.cell(r0, 1), d.cell(r0, 2)])
    for i, (k, v) in enumerate(s["by_group"]):
        d.cell(r0 + 1 + i, 1, k).font = base
        c = d.cell(r0 + 1 + i, 2, v)
        c.font, c.number_format = base, "#,##0"
    g_end = r0 + max(len(s["by_group"]), 1)

    r1 = g_end + 3
    d.cell(r1, 1, t(lang, "date")), d.cell(r1, 2, t(lang, "amount"))
    _style_header([d.cell(r1, 1), d.cell(r1, 2)])
    for i, (k, v) in enumerate(s["series"].items()):
        d.cell(r1 + 1 + i, 1, k).font = base
        c = d.cell(r1 + 1 + i, 2, v)
        c.font, c.number_format = base, "#,##0"
    s_end = r1 + max(len(s["series"]), 1)
    d.column_dimensions["A"].width = 30
    d.column_dimensions["B"].width = 16

    if rows:
        ch = BarChart()
        ch.type, ch.title, ch.legend = "bar", t(lang, "cat_chart"), None
        ch.add_data(Reference(d, min_col=2, min_row=r0, max_row=g_end), titles_from_data=True)
        ch.set_categories(Reference(d, min_col=1, min_row=r0 + 1, max_row=g_end))
        ch.height, ch.width = 9, 18
        d.add_chart(ch, "D3")
        ch2 = BarChart()
        ch2.title, ch2.legend = t(lang, s["series_label"]), None
        ch2.add_data(Reference(d, min_col=2, min_row=r1, max_row=s_end), titles_from_data=True)
        ch2.set_categories(Reference(d, min_col=1, min_row=r1 + 1, max_row=s_end))
        ch2.height, ch2.width = 8, 18
        d.add_chart(ch2, "D22")

    wb.move_sheet(det, offset=wb.sheetnames.index(t(lang, "dashboard")) + 1 - wb.sheetnames.index(det.title))


def _inject_cache(data: bytes, sheet_xml: str, values: dict) -> bytes:
    """openpyxl writes formulas without results. Add the computed results so viewers that don't
    recalculate (phone previewers) still show numbers; Excel recalculates anyway (fullCalcOnLoad)."""
    zin = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    pat = re.compile(r'<c r="([A-Z]+[0-9]+)"([^>]*)><f>(.*?)</f><v\s*/></c>|<c r="([A-Z]+[0-9]+)"([^>]*)><f>(.*?)</f><v></v></c>',
                     re.S)

    def repl(m):
        coord, attrs, f = (m.group(1), m.group(2), m.group(3)) if m.group(1) else (m.group(4), m.group(5), m.group(6))
        if coord not in values:
            return m.group(0)
        v = values[coord]
        if isinstance(v, str):
            attrs = re.sub(r'\st="[^"]*"', "", attrs) + ' t="str"'
            return f'<c r="{coord}"{attrs}><f>{f}</f><v>{escape(v)}</v></c>'
        return f'<c r="{coord}"{attrs}><f>{f}</f><v>{round(v, 6)!r}</v></c>'

    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            blob = zin.read(item.filename)
            if item.filename == sheet_xml:
                blob = pat.sub(repl, blob.decode("utf-8")).encode("utf-8")
            zout.writestr(item, blob)
    return out.getvalue()


def build_excel(rows, title, lang, period_text="", year=None):
    year = year or (int(rows[-1]["ts"][:4]) if rows else None)
    if year is None:
        from db import now
        year = now().year
    wb = load_workbook(TEMPLATE)
    cache = _localize_and_fill(wb, rows, lang, year)
    _add_dashboard_and_details(wb, rows, lang, title, period_text or title)
    wb.active = 0
    wb.calculation = CalcProperties(fullCalcOnLoad=True)
    buf = io.BytesIO()
    wb.save(buf)
    # Budget is the first sheet -> xl/worksheets/sheet1.xml
    return io.BytesIO(_inject_cache(buf.getvalue(), "xl/worksheets/sheet1.xml", cache))
